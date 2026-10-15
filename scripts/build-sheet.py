"""Sestaví soubor KOGO-texty-webu.xlsx z řádků, které vypíše scripts/sheet-rows.mjs.
Soubor se nahraje do Google Disku a otevře jako Google tabulka."""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

rows = json.load(open(sys.argv[1]))
out = sys.argv[2]
GOLD, CREAM, MUTED, LINE = "8B7041", "F4EFE6", "857B72", "E2DAD0"
head_font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
head_fill = PatternFill("solid", fgColor=GOLD)
group_fill = PatternFill("solid", fgColor=CREAM)
id_font = Font(color="B5ADA3", name="Arial", size=8)
body = Font(name="Arial", size=10)
wrap = Alignment(wrap_text=True, vertical="top")
thin = Border(bottom=Side(style="thin", color=LINE))

def sheet(wb, title, headers, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        c = ws.cell(1, i); c.font = head_font; c.fill = head_fill; c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    return ws

wb = Workbook(); wb.remove(wb.active)

# --- Návod ----------------------------------------------------------------
ws = wb.create_sheet("Návod")
ws.column_dimensions["A"].width = 110
guide = [
    ("KOGO Slovanský dům – texty webu", Font(bold=True, size=16, color=GOLD, name="Arial")),
    ("", None),
    ("Tady upravujete texty, ceny, jídelní lístek a kalendář akcí na webu. Změny se na webu objeví samy do 15 minut.", body),
    ("", None),
    ("JAK NA TO", Font(bold=True, color=GOLD, name="Arial")),
    ("1. Najděte list podle toho, co chcete změnit (záložky dole).", body),
    ("2. Přepište text přímo v buňce – česky i anglicky. Nic dalšího není potřeba.", body),
    ("3. Do 15 minut uvidíte změnu na webu.", body),
    ("", None),
    ("LISTY", Font(bold=True, color=GOLD, name="Arial")),
    ("Texty – všechny texty webu, seřazené podle stránek. Sloupec „Kde na webu“ říká, kde text najdete.", body),
    ("Kontakty a ceny – telefony, e-mail, adresa, otevírací doba, ceny menu na akce.", body),
    ("Jídelní lístek – jídla, popisy, ceny a alergeny. Jídlo přidáte novým řádkem, smažete smazáním řádku.", body),
    ("Kalendář akcí – nové akce. Co řádek, to akce. Proběhlé akce se z webu skryjí samy.", body),
    ("", None),
    ("CO NEDĚLAT", Font(bold=True, color=GOLD, name="Arial")),
    ("• Neměňte šedé sloupce „ID“ – podle nich web pozná, kam text patří.", body),
    ("• Nepřejmenovávejte listy ani nadpisy sloupců.", body),
    ("• Když buňku necháte prázdnou, na webu zůstane původní text. Nic se tím nerozbije.", body),
    ("", None),
    ("JÍDELNÍ LÍSTEK – PŘIDÁNÍ JÍDLA", Font(bold=True, color=GOLD, name="Arial")),
    ("Vložte nový řádek do správné kategorie a do šedého sloupce „ID kategorie“ zkopírujte hodnotu z řádku nad ním.", body),
    ("Cenu pište jen číslem, např. 450. Alergeny čísly oddělenými čárkou, např. 1, 3, 7.", body),
    ("", None),
    ("KALENDÁŘ AKCÍ", Font(bold=True, color=GOLD, name="Arial")),
    ("Datum pište ve tvaru 15. 11. 2026. Řádek s nápisem PŘÍKLAD je jen vzor a na web se nedostane.", body),
    ("", None),
    ("Když si nebudete jistá, napište nám – historie změn v tabulce (Soubor → Historie verzí) umí vrátit cokoli zpátky.", Font(italic=True, color=MUTED, name="Arial")),
]
for t, f in guide:
    ws.append([t])
    if f: ws.cell(ws.max_row, 1).font = f
    ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True)

# --- Texty ------------------------------------------------------------------
order = []
for t in rows["texts"]:
    if t["where"] not in order: order.append(t["where"])
ws = sheet(wb, "Texty", ["Kde na webu", "Česky", "Anglicky", "ID (neměnit)"], [30, 60, 60, 18])
last = None
for t in sorted(rows["texts"], key=lambda t: order.index(t["where"])):
    ws.append([t["where"] if t["where"] != last else "", t["cs"], t["en"], t["key"]])
    r = ws.max_row
    for c in range(1, 4): ws.cell(r, c).font = body; ws.cell(r, c).alignment = wrap
    ws.cell(r, 4).font = id_font
    if t["where"] != last:
        for c in range(1, 5): ws.cell(r, c).border = Border(top=Side(style="thin", color=GOLD))
        ws.cell(r, 1).font = Font(bold=True, name="Arial", size=10, color=GOLD)
    last = t["where"]

# --- Kontakty a ceny --------------------------------------------------------
NAMES = {
    "site.contact.addressLine": "Adresa", "site.contact.phone": "Hlavní telefon (restaurace)", "site.contact.email": "E-mail",
    "site.hours.0.time": "Otevírací doba restaurace (souhrn)", "site.venues.0.phone": "Restaurace – telefon",
    "site.venues.0.hours.0.time": "Restaurace – otevírací doba", "site.venues.1.phone": "Wine bar – telefon",
    "site.venues.1.hours.0.time": "Wine bar – otevírací doba",
    "events.eventMenus.0.prices.0.price": "Tříchodové menu – vegetariánské (cena za osobu)",
    "events.eventMenus.0.prices.1.price": "Tříchodové menu – rybí / masové (cena za osobu)",
    "events.eventMenus.1.prices.0.price": "Čtyřchodové menu – vegetariánské (cena za osobu)",
    "events.eventMenus.1.prices.1.price": "Čtyřchodové menu – rybí / masové (cena za osobu)",
    "events.eventPresentation.pricing.rows.0.three": "Prezentace eventů – ceník: vegetariánské, 3 chody",
    "events.eventPresentation.pricing.rows.0.four": "Prezentace eventů – ceník: vegetariánské, 4 chody",
    "events.eventPresentation.pricing.rows.1.three": "Prezentace eventů – ceník: rybí / masové, 3 chody",
    "events.eventPresentation.pricing.rows.1.four": "Prezentace eventů – ceník: rybí / masové, 4 chody",
}
ws = sheet(wb, "Kontakty a ceny", ["Co", "Hodnota", "ID (neměnit)"], [52, 40, 18])
for v in rows["values"]:
    ws.append([NAMES.get(v["key"], v["where"]), v["value"], v["key"]])
    r = ws.max_row; ws.cell(r, 1).font = body; ws.cell(r, 2).font = body; ws.cell(r, 3).font = id_font

# --- Jídelní lístek ---------------------------------------------------------
ws = sheet(wb, "Jídelní lístek", ["Kategorie", "Název jídla", "Popis česky", "Popis anglicky", "Cena (Kč)", "Alergeny", "ID kategorie (neměnit)"], [22, 34, 44, 44, 11, 13, 18])
last = None
for m in rows["menu"]:
    ws.append([m["section"] if m["id"] != last else "", m["name"], m["cs"], m["en"], str(m["price"]) if m["price"] is not None else "", m["allergens"], m["id"]])
    r = ws.max_row
    for c in range(1, 7): ws.cell(r, c).font = body; ws.cell(r, c).alignment = wrap
    ws.cell(r, 7).font = id_font
    if m["id"] != last:
        for c in range(1, 8): ws.cell(r, c).border = Border(top=Side(style="thin", color=GOLD))
        ws.cell(r, 1).font = Font(bold=True, name="Arial", size=10, color=GOLD)
    last = m["id"]

# --- Kalendář akcí ----------------------------------------------------------
ws = sheet(wb, "Kalendář akcí", ["Datum", "Čas", "Název česky", "Název anglicky", "Popis česky", "Popis anglicky", "Odkaz (nepovinné)"], [14, 14, 30, 30, 50, 50, 30])
ws.append(["PŘÍKLAD", "18:00", "Svatomartinské hody", "St. Martin's feast", "Tradiční husa se zelím a knedlíkem…", "Traditional goose with cabbage…", ""])
for c in range(1, 8): ws.cell(2, c).font = Font(italic=True, color=MUTED, name="Arial", size=10); ws.cell(2, c).fill = group_fill

# Všechny buňky jako prostý text – jinak Google převádí data a čísla a export
# listu by buňky jiného typu než většina sloupce vynechal. Do kalendáře a lístku
# rovnou připravíme prázdné řádky ve stejném formátu.
for ws in wb.worksheets[1:]:
    extra = 200 if ws.title == "Kalendář akcí" else 60
    for r in range(2, ws.max_row + extra):
        for c in range(1, ws.max_column + 1):
            ws.cell(r, c).number_format = "@"

wb.active = 0
wb.save(out)
print("uloženo:", out)
